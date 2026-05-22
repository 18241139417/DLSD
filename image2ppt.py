#!/usr/bin/env python3
"""
image2ppt.py

Usage examples:
  python image2ppt.py --input slide.png --output slide.pptx
  python image2ppt.py --input mockup.html --output mockup.pptx
  python image2ppt.py --input diagram.svg --output diagram.pptx
"""

from __future__ import annotations

import argparse
import base64
import io
import json
import os
import re
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

import anthropic
import cairosvg
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from playwright.sync_api import sync_playwright


SLIDE_WIDTH_INCH = 13.33
SLIDE_HEIGHT_INCH = 7.5
TARGET_WIDTH = 1920
TARGET_HEIGHT = 1080
MODEL_NAME = "claude-sonnet-4-20250514"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert PNG/JPEG/HTML/SVG into editable PPTX via Claude vision layout extraction."
    )
    parser.add_argument("--input", required=True, help="Input file path: png/jpeg/html/svg")
    parser.add_argument("--output", required=True, help="Output .pptx file path")
    return parser.parse_args()


def normalize_to_png(input_path: Path) -> bytes:
    suffix = input_path.suffix.lower()

    if suffix in {".png", ".jpg", ".jpeg"}:
        return image_file_to_png(input_path)
    if suffix in {".html", ".htm"}:
        return html_to_png(input_path)
    if suffix == ".svg":
        return svg_to_png(input_path)

    raise ValueError(f"Unsupported input type: {suffix}")


def image_file_to_png(path: Path) -> bytes:
    with Image.open(path) as img:
        if img.mode != "RGB":
            img = img.convert("RGB")

        canvas = Image.new("RGB", (TARGET_WIDTH, TARGET_HEIGHT), color=(255, 255, 255))
        src_w, src_h = img.size
        scale = min(TARGET_WIDTH / src_w, TARGET_HEIGHT / src_h)
        new_size = (max(1, int(src_w * scale)), max(1, int(src_h * scale)))
        resized = img.resize(new_size, Image.Resampling.LANCZOS)
        offset_x = (TARGET_WIDTH - new_size[0]) // 2
        offset_y = (TARGET_HEIGHT - new_size[1]) // 2
        canvas.paste(resized, (offset_x, offset_y))

        buf = io.BytesIO()
        canvas.save(buf, format="PNG")
        return buf.getvalue()


def html_to_png(path: Path) -> bytes:
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": TARGET_WIDTH, "height": TARGET_HEIGHT})
        page.goto(path.resolve().as_uri(), wait_until="networkidle")
        image = page.screenshot(type="png", full_page=False)
        browser.close()
        return image


def svg_to_png(path: Path) -> bytes:
    return cairosvg.svg2png(
        url=str(path),
        output_width=TARGET_WIDTH,
        output_height=TARGET_HEIGHT,
    )


def get_anthropic_client() -> anthropic.Anthropic:
    api_key = os.getenv("ANTHROPIC_API_KEY")
    if not api_key:
        raise EnvironmentError("ANTHROPIC_API_KEY environment variable is required")
    return anthropic.Anthropic(api_key=api_key)


def analyze_layout_with_claude(client: anthropic.Anthropic, png_data: bytes) -> Dict[str, Any]:
    prompt = (
        "You are an assistant that extracts slide layouts from an image. "
        "Return only pure JSON (no markdown fences, no comments). "
        "Schema:\n"
        "{\n"
        "  \"background_color\": \"#RRGGBB\",\n"
        "  \"elements\": [\n"
        "    {\n"
        "      \"type\": \"text\" | \"shape\",\n"
        "      \"role\": \"title\" | \"subtitle\" | \"body\" | \"caption\" | \"decoration\" | \"other\",\n"
        "      \"text\": string,\n"
        "      \"x\": number,\n"
        "      \"y\": number,\n"
        "      \"width\": number,\n"
        "      \"height\": number,\n"
        "      \"font_size\": number,\n"
        "      \"font_bold\": boolean,\n"
        "      \"font_color\": \"#RRGGBB\",\n"
        "      \"align\": \"left\" | \"center\" | \"right\",\n"
        "      \"fill_color\": \"#RRGGBB\",\n"
        "      \"border_color\": \"#RRGGBB\"\n"
        "    }\n"
        "  ]\n"
        "}\n"
        "All coordinates and sizes must be relative values between 0 and 1. "
        "For missing values, provide reasonable defaults."
    )

    resp = client.messages.create(
        model=MODEL_NAME,
        max_tokens=4000,
        temperature=0,
        messages=[
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt},
                    {
                        "type": "image",
                        "source": {
                            "type": "base64",
                            "media_type": "image/png",
                            "data": base64.b64encode(png_data).decode("utf-8"),
                        },
                    },
                ],
            }
        ],
    )

    text = "".join(block.text for block in resp.content if getattr(block, "type", "") == "text")
    return parse_json_from_text(text)


def parse_json_from_text(text: str) -> Dict[str, Any]:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)

    try:
        return json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, flags=re.DOTALL)
        if not match:
            raise
        return json.loads(match.group(0))


def clamp01(v: Any, default: float) -> float:
    try:
        x = float(v)
    except (TypeError, ValueError):
        x = default
    return max(0.0, min(1.0, x))


def parse_hex_color(value: Optional[str], default: str = "#000000") -> RGBColor:
    s = (value or default).strip()
    if not re.fullmatch(r"#?[0-9a-fA-F]{6}", s):
        s = default
    s = s.lstrip("#")
    return RGBColor(int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))


def map_align(align: str) -> PP_ALIGN:
    align = (align or "left").lower()
    if align == "center":
        return PP_ALIGN.CENTER
    if align == "right":
        return PP_ALIGN.RIGHT
    return PP_ALIGN.LEFT


def build_ppt(layout: Dict[str, Any], output_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH_INCH)
    prs.slide_height = Inches(SLIDE_HEIGHT_INCH)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = parse_hex_color(layout.get("background_color"), "#FFFFFF")

    elements: List[Dict[str, Any]] = layout.get("elements", [])

    for el in elements:
        typ = (el.get("type") or "").lower()
        left = Inches(SLIDE_WIDTH_INCH * clamp01(el.get("x"), 0.0))
        top = Inches(SLIDE_HEIGHT_INCH * clamp01(el.get("y"), 0.0))
        width = Inches(SLIDE_WIDTH_INCH * clamp01(el.get("width"), 0.2))
        height = Inches(SLIDE_HEIGHT_INCH * clamp01(el.get("height"), 0.1))

        if typ == "text":
            shape = slide.shapes.add_textbox(left, top, width, height)
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.alignment = map_align(el.get("align", "left"))
            run = p.add_run()
            run.text = str(el.get("text") or "")
            run.font.size = Pt(float(el.get("font_size") or 18))
            run.font.bold = bool(el.get("font_bold", False))
            run.font.color.rgb = parse_hex_color(el.get("font_color"), "#000000")

            fill_color = el.get("fill_color")
            if fill_color:
                shape.fill.solid()
                shape.fill.fore_color.rgb = parse_hex_color(fill_color)
            else:
                shape.fill.background()

            border_color = el.get("border_color")
            if border_color:
                shape.line.color.rgb = parse_hex_color(border_color)
            else:
                shape.line.fill.background()

        elif typ == "shape":
            from pptx.enum.shapes import MSO_SHAPE

            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
            shape.fill.solid()
            shape.fill.fore_color.rgb = parse_hex_color(el.get("fill_color"), "#DDDDDD")
            border_color = el.get("border_color")
            if border_color:
                shape.line.color.rgb = parse_hex_color(border_color)
            else:
                shape.line.fill.background()

    prs.save(str(output_path))


def main() -> int:
    args = parse_args()
    input_path = Path(args.input)
    output_path = Path(args.output)

    if not input_path.exists():
        print(f"Input not found: {input_path}", file=sys.stderr)
        return 1

    try:
        png_data = normalize_to_png(input_path)
        client = get_anthropic_client()
        layout = analyze_layout_with_claude(client, png_data)
        build_ppt(layout, output_path)
    except Exception as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1

    print(f"Saved PPTX: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
