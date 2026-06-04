#!/usr/bin/env python3
import argparse, json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

SHAPE_MAP = {
    'rect': MSO_AUTO_SHAPE_TYPE.RECTANGLE,
    'roundRect': MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
    'ellipse': MSO_AUTO_SHAPE_TYPE.OVAL,
}

def hex_to_rgb(v):
    v=v.lstrip('#')
    return RGBColor(int(v[0:2],16), int(v[2:4],16), int(v[4:6],16))

def px_to_in(px, canvas, slide_in):
    return (px / canvas) * slide_in

def build(manifest_path, output_path, summary_path):
    data = json.loads(Path(manifest_path).read_text())
    deck = data['deck']
    cw, ch = deck['canvas_width'], deck['canvas_height']
    sw = float(deck.get('slide_width_in', 13.333))
    sh = sw * ch / cw
    prs = Presentation()
    prs.slide_width = Inches(sw)
    prs.slide_height = Inches(sh)

    counts = {'slides':0,'text':0,'images':0,'shapes':0,'backgrounds':0}

    blank = prs.slide_layouts[6]
    base_dir = Path(manifest_path).parent
    for s in data['slides']:
        slide = prs.slides.add_slide(blank)
        counts['slides'] += 1
        for el in s.get('elements',[]):
            kind = el['kind']
            if kind == 'background':
                counts['backgrounds'] += 1
                if 'fill' in el:
                    left=top=Inches(0)
                    w,h = Inches(sw), Inches(sh)
                    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,left,top,w,h)
                    shp.fill.solid(); shp.fill.fore_color.rgb = hex_to_rgb(el['fill'])
                    shp.line.fill.background()
                elif 'file' in el:
                    p = (base_dir/el['file']).resolve()
                    slide.shapes.add_picture(str(p), Inches(0), Inches(0), width=Inches(sw), height=Inches(sh))
                    counts['images'] += 1
            elif kind == 'image':
                p = (base_dir/el['file']).resolve()
                x=Inches(px_to_in(el['x'],cw,sw)); y=Inches(px_to_in(el['y'],ch,sh))
                w=Inches(px_to_in(el['w'],cw,sw)); h=Inches(px_to_in(el['h'],ch,sh))
                slide.shapes.add_picture(str(p), x,y,width=w,height=h)
                counts['images'] += 1
            elif kind == 'text':
                x=Inches(px_to_in(el['x'],cw,sw)); y=Inches(px_to_in(el['y'],ch,sh))
                w=Inches(px_to_in(el['w'],cw,sw)); h=Inches(px_to_in(el['h'],ch,sh))
                tx = slide.shapes.add_textbox(x,y,w,h)
                tf = tx.text_frame
                tf.clear()
                p = tf.paragraphs[0]
                p.text = el.get('text','')
                run = p.runs[0]
                run.font.name = el.get('font_family','Arial')
                if 'font_size_pt' in el: run.font.size = Pt(el['font_size_pt'])
                elif 'font_size_px' in el: run.font.size = Pt(el['font_size_px']*0.75)
                run.font.bold = bool(el.get('bold',False))
                if 'color' in el: run.font.color.rgb = hex_to_rgb(el['color'])
                counts['text'] += 1
            elif kind == 'shape':
                st = SHAPE_MAP.get(el.get('shape','rect'), MSO_AUTO_SHAPE_TYPE.RECTANGLE)
                x=Inches(px_to_in(el['x'],cw,sw)); y=Inches(px_to_in(el['y'],ch,sh))
                w=Inches(px_to_in(el['w'],cw,sw)); h=Inches(px_to_in(el['h'],ch,sh))
                shp = slide.shapes.add_shape(st,x,y,w,h)
                if 'fill' in el:
                    shp.fill.solid(); shp.fill.fore_color.rgb = hex_to_rgb(el['fill'])
                if 'stroke' in el:
                    shp.line.color.rgb = hex_to_rgb(el['stroke'])
                counts['shapes'] += 1

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    Path(summary_path).write_text(json.dumps({'counts':counts,'slide_width_in':sw,'slide_height_in':sh},indent=2))

def main():
    ap=argparse.ArgumentParser()
    sp=ap.add_subparsers(dest='cmd', required=True)
    b=sp.add_parser('build')
    b.add_argument('--manifest', required=True)
    b.add_argument('--output', required=True)
    b.add_argument('--summary', required=True)
    args=ap.parse_args()
    if args.cmd=='build': build(args.manifest,args.output,args.summary)

if __name__=='__main__': main()
